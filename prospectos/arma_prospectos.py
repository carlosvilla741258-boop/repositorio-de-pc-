"""Arma el Excel de prospectos (negocios de SJL sin teléfono: ese dato lo llena el usuario)."""
import sys
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = sys.argv[1]

F = "Arial"
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
CABECERA = PatternFill("solid", fgColor="1F4E5A")
ZEBRA = PatternFill("solid", fgColor="F3F6F7")
LINEA = Side(style="thin", color="D0D7DA")
BORDE = Border(bottom=LINEA)

Z1 = "Av. Las Flores de Primavera"
Z2 = "Av. Los Jardines"
Z3 = "Próceres: Las Flores / La Hacienda"
Z4 = "Los Postes / San Hilarión"
Z5 = "San Carlos"
ZONAS = [Z1, Z2, Z3, Z4, Z5]

IDEA = {
    "Dentista": "Web con tratamientos y precios referenciales + botón de cita por WhatsApp",
    "Veterinaria": "Web con servicios, horario de emergencias y citas por WhatsApp",
    "Belleza": "Galería de trabajos, lista de precios y reservas por WhatsApp",
    "Gimnasio": "Planes, horarios y fotos del local + inscripción por WhatsApp",
    "Comida": "Carta digital con QR + pedidos de delivery por WhatsApp",
    "Hostal": "Fotos de habitaciones, tarifas y reservas por WhatsApp; salir en Google",
    "Educación": "Niveles o ciclos, admisión 2027, galería y formulario de inscripción",
    "Notaría": "Trámites con requisitos, horario y cotización por WhatsApp",
    "Taller": "Servicios, fotos de antes/después y garantía (usa tu web de Abarca Factoría como ejemplo)",
    "Tienda": "Catálogo en línea con pedidos por WhatsApp",
}
MEJORAR = "Ya tiene web: revísala en el celular y ofrécele rediseño, dominio propio o botón de WhatsApp"

AQUI = "https://aquihoteles.com/peru/lima/lima/san-juan-de-lurigancho/"
BUSQ = None  # vino solo del resumen de una búsqueda web, sin página propia que citar

# (negocio, rubro, tipo, zona, dirección, referencia, web, oportunidad, fuente, consulta_maps)
N = [
    ("Clínica Dental Salvadent", "Dentista", "Dentista", Z1, "Av. Las Flores de Primavera 1442", "",
     "No encontrada", "Web nueva",
     "https://pe.near-place.com/clinica-dental-salvadent-avenue-flores-de-primavera-1442-san-juan-de-lurigancho/en", None),
    ("Clínica Veterinaria Maskotopía", "Veterinaria", "Veterinaria", Z1, "Av. Las Flores de Primavera 1701", "",
     "Página gratis de Wix (maskotopia.wixsite.com)", "Mejorar web",
     "https://maskotopia.wixsite.com/maskotopia", None),
    ("Veterinaria Peques Pets", "Veterinaria", "Veterinaria", Z1, "Av. Las Flores de Primavera 1117", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Mundo Pets (sede Las Flores)", "Veterinaria", "Veterinaria", Z1, "Av. Las Flores de Primavera 1568", "",
     "Sí: veterinariamundopets.com", "Mejorar web", "https://veterinariamundopets.com/", None),
    ("Clínica Veterinaria Aristocat", "Veterinaria", "Veterinaria", Z1, "Calle Los Huertos 129, Urb. Las Flores de Lima", "",
     "No encontrada", "Web nueva",
     "https://www.cylex.com.pe/san-juan-de-lurigancho/clinica-veterinaria-aristocat-11167731.html", None),
    ("Gym Mega Force (sede Las Flores)", "Gimnasio", "Gimnasio", Z1, "Av. Las Flores de Primavera 1238",
     "Misma marca que la sede Hacienda (fila de abajo): una web para las 2 sedes",
     "Solo redes (Facebook)", "Web nueva",
     "https://feelingperu.com/gym-mega-force-las-flores-san-juan-de-lurigancho/", None),
    ("El Capitán Barbería", "Barbería", "Belleza", Z1, "Av. Las Flores de Primavera 1255",
     "Paradero 8, frente a la discoteca Kurda; a minutos de la Estación Los Jardines",
     "Solo redes (TikTok)", "Web nueva",
     "https://www.tiktok.com/@elcapitanbarberia/video/7392704411308526854", None),
    ("El Legado Barbería", "Barbería", "Belleza", Z1, "Av. Las Flores de Primavera 1285, Urb. Las Flores", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Punto 8 Barbería SJL", "Barbería", "Belleza", Z1, "Av. Las Flores de Primavera 1211", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Cambio de Look", "Peluquería", "Belleza", Z1, "Av. Las Flores de Primavera 1089", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Luna High Studio", "Salón de belleza y cafetería", "Belleza", Z1, "Av. Las Flores de Primavera 1301", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Corralito – Pollos y Parrillas", "Pollería", "Comida", Z1, "Av. Las Flores de Primavera 1787", "",
     "No encontrada", "Web nueva",
     "https://restaurantguru.com/Polleria-Corralito-San-Juan-de-Lurigancho-2", None),
    ("Pollería Ricco", "Pollería", "Comida", Z1, "Av. Las Flores de Primavera (número por confirmar)",
     "El directorio no da el número de la puerta: ubícala en Maps",
     "No encontrada", "Web nueva",
     "https://restaurantguru.com/Polleria-Ricco-San-Juan-de-Lurigancho",
     "Pollería Ricco, Av. Las Flores de Primavera, San Juan de Lurigancho, Lima"),
    ("Las Flores MTL", "Restaurante con pedidos en línea", "Comida", Z1, "Av. Las Flores de Primavera 1163", "",
     "Tienda en ola.click (plataforma de pedidos)", "Mejorar web",
     "https://mi-tercer-lugar-san-juan-de-lurigancho.ola.click/products", None),
    ("Colegio Jean Piaget – Las Flores", "Colegio", "Educación", Z1, "Jr. Milenramas 752, Urb. Las Flores de Primavera", "",
     "Solo redes (Facebook)", "Web nueva",
     "https://www.facebook.com/ColegioJeanPiagetLasFlores/", None),

    ("Hostal La Hacienda", "Hostal", "Hostal", Z2, "Av. Los Jardines Oeste 247", "",
     "No encontrada", "Web nueva", "https://es.cybo.com/PE-biz/hostal-la-hacienda", None),
    ("Odonto Marcos", "Dentista", "Dentista", Z2, "Av. Los Jardines Oeste 249", "",
     "Solo redes (Instagram)", "Web nueva", "https://www.instagram.com/clinicaodontomarcos.sjl/", None),
    ("Clínica Dental Implant SJL", "Dentista", "Dentista", Z2, "Av. Los Jardines Oeste 160, 2.º piso", "",
     "Sí: clinicadentalimplant.com", "Mejorar web", "https://clinicadentalimplant.com/", None),
    ("AMALY Spa", "Spa y uñas", "Belleza", Z2, "Av. Los Jardines Oeste 125", "",
     "Solo redes (Instagram)", "Web nueva", "https://www.instagram.com/amaly.spa/", None),
    ("Salón & Spa Venus", "Salón y spa", "Belleza", Z2, "Av. Los Jardines Oeste 239A, Urb. Las Flores de Lima", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Vanidosas Salon & Spa", "Salón y spa", "Belleza", Z2, "Av. Los Jardines Este 283, Galería Todos, local 117", "",
     "Solo redes (Facebook)", "Web nueva", "https://www.facebook.com/Vanidosas.salonspa/",
     "Vanidosas Salon & Spa, Av. Los Jardines Este 283, San Juan de Lurigancho, Lima"),
    ("Chifa Los Jardines", "Chifa", "Comida", Z2, "Av. Los Jardines Este 176", "",
     "No encontrada", "Web nueva",
     "https://www.waze.com/live-map/directions/chifa-los-jardines-av.-los-jardines-este-176-san-juan-de-lurigancho?to=place.w.185468560.1854620063.16915909",
     None),
    ("Restaurante Cevichería El Buen Sabor", "Cevichería", "Comida", Z2, "Av. Los Jardines Este 139", "",
     "No encontrada", "Web nueva",
     "https://carta.menu/restaurants/san-juan-de-lurigancho/restaurante-cevicheria-el-buen-sabor", None),
    ("Cevichería Don Pez", "Cevichería", "Comida", Z2, "Jr. Los Ricinos 1039, Los Jardines", "",
     "No encontrada", "Web nueva",
     "https://restaurantguru.com/Restaurante-Cevicheria-Don-Pez-San-Juan-de-Lurigancho", None),
    ("Veterinaria Posta Oasis (Jardines)", "Veterinaria", "Veterinaria", Z2, "Av. 13 de Enero 1708", "",
     "Solo redes (Facebook)", "Web nueva", BUSQ, None),

    ("Chifa Yao Fu", "Chifa", "Comida", Z3, "Av. Próceres de la Independencia 728", "",
     "No encontrada", "Web nueva", "https://www.cybo.com/PE-biz/chifa-yao-fu_2b", None),
    ("Hostal Jerusalén", "Hostal", "Hostal", Z3, "Av. Próceres de la Independencia 595, Urb. Las Flores", "",
     "No encontrada", "Web nueva", AQUI, None),
    ("Hostal Cielo", "Hostal", "Hostal", Z3, "Av. Próceres de la Independencia 576", "Altura de la Posta San Juan",
     "No encontrada", "Web nueva", AQUI, None),
    ("Gym Mega Force (sede Hacienda)", "Gimnasio", "Gimnasio", Z3, "Av. Próceres de la Independencia 1715",
     "Misma marca que la sede Las Flores", "Solo redes (Facebook)", "Web nueva",
     "https://www.facebook.com/megaforce.hacienda/", None),
    ("Notaría Villota", "Notaría", "Notaría", Z3, "Av. Próceres de la Independencia 1722-B, of. 204, Urb. Las Flores", "",
     "No encontrada", "Web nueva",
     "https://www.emis.com/php/company-profile/PE/Notaria_Villota__villota_Cerna_Marco_Antonio__es_9666830.html",
     "Notaría Villota, Av. Próceres de la Independencia 1722, San Juan de Lurigancho, Lima"),

    ("Hostal Nubeluz", "Hostal", "Hostal", Z4, "Av. Próceres de la Independencia 2001, Urb. San Hilarión", "",
     "No encontrada", "Web nueva", AQUI, None),
    ("Hostal Castillo", "Hostal", "Hostal", Z4, "Av. Los Postes 233, Asoc. Las Begonias Mz. D Lt. 20", "",
     "No encontrada", "Web nueva",
     "https://aquihoteles.com/peru/lima/lima/san-juan-de-lurigancho/hostal/hostal-en-san-juan-de-lurigancho-castillo/",
     "Hostal Castillo, Av. Los Postes 233, San Juan de Lurigancho, Lima"),
    ("Clínica Dental Odontomanía", "Dentista", "Dentista", Z4, "Av. Próceres de la Independencia 2286, 2.º piso", "",
     "Solo redes (TikTok)", "Web nueva",
     "https://www.tiktok.com/@odontomania/video/7380515548674755846", None),
    ("La Pollerona – San Hilarión", "Pollería", "Comida", Z4, "Av. Próceres de la Independencia 2292", "",
     "No encontrada", "Web nueva",
     "https://wanderboat.ai/restaurants/peru/lima-metropolitan-area/la-pollerona-san-hilarion/9i3EAUCeSNGaXDyUKQR4bQ",
     None),
    ("SAVIA Academia Preuniversitaria", "Academia preuniversitaria", "Educación", Z4,
     "Av. Próceres de la Independencia 2175", "", "No encontrada", "Web nueva", BUSQ, None),
    ("A la Gloria Jugos y Sándwiches", "Juguería", "Comida", Z4, "Av. Los Postes Este 346",
     "Tiene otro local en Av. El Sol 498 (frente a la UPN)", "Sí: alagloria.com", "Mejorar web",
     "https://restaurantguru.com/SANDWICH-Y-JUGOS-A-LA-GLORIA-San-Juan-de-Lurigancho", None),

    ("Restaurante Moche – San Carlos", "Restaurante", "Comida", Z5, "Av. Próceres de la Independencia 2721",
     "Muy concurrido: más de 400 reseñas en Restaurant Guru", "No encontrada", "Web nueva",
     "https://restaurantguru.com/Restaurant-Moche-San-Juan-de-Lurigancho", None),
    ("San Carlos Automotriz", "Taller automotriz (GNV/GLP)", "Taller", Z5, "Av. Próceres de la Independencia 2735", "",
     "Solo redes (Facebook)", "Web nueva",
     "https://www.paginasamarillas.com.pe/empresas/san-carlos-automotriz/san-juan-de-lurigancho-584634", None),
    ("Tienda Kawaii", "Tienda de accesorios", "Tienda", Z5, "Av. El Sol 252-254, Urb. San Carlos", "Frente a la UTP",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Cevichería d'Adrián", "Cevichería", "Comida", Z5, "Av. El Sol 466", "",
     "No encontrada", "Web nueva", BUSQ, None),
    ("Hostal El Sol", "Hostal", "Hostal", Z5, "Av. El Sol Mz. D-1 Lt. 15 (N.° 511), Urb. San Carlos",
     "Altura del paradero San Carlos", "No encontrada", "Web nueva",
     "https://aquihoteles.com/peru/lima/lima/san-juan-de-lurigancho/hostal/hostal-en-san-juan-de-lurigancho-el-sol/",
     "Hostal El Sol, Av. El Sol 511, San Juan de Lurigancho, Lima"),
    ("Hostal Santa Cruz", "Hostal", "Hostal", Z5, "Jr. Las Rocas 2424, Urb. San Carlos", "",
     "No encontrada", "Web nueva", AQUI, None),
    ("Restaurant El Norteño", "Restaurante", "Comida", Z5, "Av. El Sol y Av. Wiesse s/n, int. T-29, Urb. San Carlos", "",
     "No encontrada", "Web nueva", "https://deperu.com/comercios/restaurantes/restaurant-el-norteno-57153",
     "Restaurant El Norteño, Av. El Sol, Urb. San Carlos, San Juan de Lurigancho, Lima"),
]

ESTADOS = ["Pendiente", "Contactado", "Interesado", "Cita / visita", "Vendido", "No interesado"]


def maps(consulta):
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(consulta)


def fuente_celda(celda, url):
    if url:
        celda.value = "Ver fuente"
        celda.hyperlink = url
        celda.font = Font(name=F, size=10, color="0563C1", underline="single")
    else:
        celda.value = "Búsqueda web"
        celda.font = Font(name=F, size=10, color="6B7478")


wb = Workbook()

# --- Hoja 1: Negocios ---------------------------------------------------------
ws = wb.active
ws.title = "Negocios"
ws["A1"] = ("Prospectos para páginas web · San Juan de Lurigancho: La Hacienda, Av. Las Flores, "
            "Av. Los Jardines, Los Postes y San Carlos")
ws["A1"].font = Font(name=F, size=14, bold=True, color="1F4E5A")
ws["A2"] = (f"{len(N)} negocios sacados de directorios públicos (septiembre 2026). Toca «Abrir en Google Maps»: "
            "ahí sale el teléfono y la ubicación exacta. Las columnas amarillas las llenas tú.")
ws["A2"].font = Font(name=F, size=10, italic=True, color="44535A")

CAB = ["N.°", "Negocio", "Rubro", "Zona", "Dirección", "Referencia", "Ubicación", "¿Tiene web?",
       "Oportunidad", "Qué ofrecerle", "Fuente", "Teléfono (lo llenas tú)", "Estado",
       "Fecha de contacto", "Notas"]
ANCHO = [5, 32, 22, 24, 42, 36, 22, 30, 14, 48, 13, 18, 15, 13, 34]
FILA_CAB = 4
for i, (texto, ancho) in enumerate(zip(CAB, ANCHO), start=1):
    c = ws.cell(row=FILA_CAB, column=i, value=texto)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = CABECERA
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = ancho
ws.row_dimensions[FILA_CAB].height = 30

primera = FILA_CAB + 1
for k, (negocio, rubro, tipo, zona, direccion, ref, web, oport, fuente, consulta) in enumerate(N):
    r = primera + k
    consulta = consulta or f"{negocio}, {direccion}, San Juan de Lurigancho, Lima"
    fila = [k + 1, negocio, rubro, zona, direccion, ref, None, web, oport,
            MEJORAR if oport == "Mejorar web" else IDEA[tipo], None, None, "Pendiente", None, None]
    for col, valor in enumerate(fila, start=1):
        c = ws.cell(row=r, column=col, value=valor)
        c.font = Font(name=F, size=10, bold=(col == 2))
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDE
        if k % 2:
            c.fill = ZEBRA
    m = ws.cell(row=r, column=7, value="Abrir en Google Maps")
    m.hyperlink = maps(consulta)
    m.font = Font(name=F, size=10, color="0563C1", underline="single")
    fuente_celda(ws.cell(row=r, column=11), fuente)
    for col in (12, 13, 14, 15):
        ws.cell(row=r, column=col).fill = AMARILLO
    ws.cell(row=r, column=14).number_format = "dd/mm/yyyy"

ultima = primera + len(N) - 1
ws.freeze_panes = ws.cell(row=primera, column=3)
ws.auto_filter.ref = f"A{FILA_CAB}:O{ultima}"

dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True,
                    showErrorMessage=True, errorTitle="Estado",
                    error="Elige un estado de la lista.")
ws.add_data_validation(dv)
dv.add(f"M{primera}:M{primera + 500}")

# --- Hoja 2: Resumen (fórmulas: se actualiza sola al cambiar el Estado) --------
rs = wb.create_sheet("Resumen")
rs.column_dimensions["A"].width = 38
rs.column_dimensions["B"].width = 12
rango = lambda col: f"Negocios!${col}${primera}:${col}${primera + 500}"


def tabla(fila, titulo, etiquetas, col):
    rs.cell(row=fila, column=1, value=titulo).font = Font(name=F, size=11, bold=True, color="1F4E5A")
    rs.cell(row=fila + 1, column=1, value="").font = Font(name=F)
    h1 = rs.cell(row=fila + 1, column=1, value="")
    h1.fill = CABECERA
    h2 = rs.cell(row=fila + 1, column=2, value="Negocios")
    h2.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    h2.fill = CABECERA
    for j, etq in enumerate(etiquetas):
        rr = fila + 2 + j
        rs.cell(row=rr, column=1, value=etq).font = Font(name=F, size=10)
        rs.cell(row=rr, column=2, value=f"=COUNTIF({rango(col)},A{rr})").font = Font(name=F, size=10)
    tot = fila + 2 + len(etiquetas)
    rs.cell(row=tot, column=1, value="Total").font = Font(name=F, size=10, bold=True)
    rs.cell(row=tot, column=2, value=f"=SUM(B{fila + 2}:B{tot - 1})").font = Font(name=F, size=10, bold=True)
    return tot + 2


rs["A1"] = "Resumen de prospectos"
rs["A1"].font = Font(name=F, size=14, bold=True, color="1F4E5A")
siguiente = tabla(3, "Por zona", ZONAS, "D")
siguiente = tabla(siguiente, "Por oportunidad", ["Web nueva", "Mejorar web"], "I")
tabla(siguiente, "Por estado (cambia al actualizar la columna Estado)", ESTADOS, "M")

# --- Hoja 3: Cómo usar --------------------------------------------------------
cu = wb.create_sheet("Cómo usar")
cu.column_dimensions["A"].width = 110
LINEAS = [
    ("Cómo usar esta lista", "titulo"),
    ("", None),
    ("1. Teléfono y ubicación exacta", "sub"),
    ("Toca «Abrir en Google Maps» en la hoja Negocios. En la ficha del negocio aparecen su teléfono, "
     "su ubicación en el mapa, las fotos del local y si ya tiene página web.", None),
    ("Si Maps no lo encuentra, toca «Ver fuente»: es la página del directorio donde apareció.", None),
    ("", None),
    ("2. Columnas amarillas: las llenas tú", "sub"),
    ("Teléfono · Estado (lista desplegable) · Fecha de contacto · Notas. La hoja Resumen cuenta sola "
     "cuántos llevas en cada estado.", None),
    ("Ejemplo de cómo llenarlas:  Teléfono «9XX XXX XXX» · Estado «Interesado» · Fecha «02/10/2026» · "
     "Notas «Pidió precio. Volver el lunes con propuesta».", None),
    ("Estados: Pendiente → Contactado → Interesado → Cita / visita → Vendido (o No interesado).", None),
    ("", None),
    ("3. Por dónde empezar", "sub"),
    ("Primero los que dicen «Web nueva» y «Solo redes»: ya publican en Facebook o TikTok, así que les "
     "importa conseguir clientes por internet, pero no tienen web. Los de «Mejorar web» ya pagaron una "
     "vez: revisa su web en el celular antes de escribirles y diles qué le falta.", None),
    ("Rubros que más pagan por una web: dentistas, veterinarias, hostales, gimnasios, academias y colegios, "
     "notarías.", None),
    ("", None),
    ("4. Mensaje sugerido (WhatsApp o en persona)", "sub"),
    ("Hola, buenas tardes. Soy [tu nombre], vecino de San Juan de Lurigancho. Hago páginas web para "
     "negocios de la zona. Vi [nombre del negocio] en [avenida] y creo que una web le ayudaría a que más "
     "clientes lo encuentren en Google y le escriban directo por WhatsApp. Le puedo mostrar una que hice "
     "para un taller de aquí mismo de SJL: [link de tu web de ejemplo]. ¿Le paso una propuesta sin "
     "compromiso?", "cita"),
    ("Si te dicen que no, no insistas: agradece y pasa al siguiente. Un «no» educado hoy puede ser un "
     "cliente dentro de unos meses.", None),
    ("", None),
    ("5. Qué NO está en la lista", "sub"),
    ("Cadenas grandes (InkaFarma, Tambo, BCP, Norky's, Roky's, Smart Fit, Ópticas GMO, COA, Multilab, "
     "Laser Dent Kids…): su web la maneja la oficina central, no el local.", None),
    ("Los datos vienen de directorios públicos y pueden estar desactualizados: confirma en Maps que el "
     "negocio sigue abierto antes de ir.", None),
]
for i, (texto, estilo) in enumerate(LINEAS, start=1):
    c = cu.cell(row=i, column=1, value=texto)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if estilo == "titulo":
        c.font = Font(name=F, size=14, bold=True, color="1F4E5A")
    elif estilo == "sub":
        c.font = Font(name=F, size=11, bold=True, color="1F4E5A")
    elif estilo == "cita":
        c.font = Font(name=F, size=10, italic=True)
        c.fill = AMARILLO
    else:
        c.font = Font(name=F, size=10)

wb.save(SALIDA)
print(f"{len(N)} negocios -> {SALIDA}")
